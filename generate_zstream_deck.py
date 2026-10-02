"""Generate the Z-Stream Test Generator overview slide deck as .pptx"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

# Colors — dark theme matching existing decks
BG_DARK = RGBColor(0x1A, 0x1A, 0x2E)
BG_MID = RGBColor(0x16, 0x21, 0x3E)
ACCENT_RED = RGBColor(0xE9, 0x45, 0x60)
ACCENT_BLUE = RGBColor(0x7E, 0xC8, 0xE3)
ACCENT_GREEN = RGBColor(0x5E, 0xC4, 0x6B)
ACCENT_YELLOW = RGBColor(0xF5, 0xC5, 0x42)
ACCENT_ORANGE = RGBColor(0xF0, 0x8A, 0x40)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xD0, 0xD0, 0xE0)
DIM_GRAY = RGBColor(0x70, 0x70, 0x90)
TABLE_BG = RGBColor(0x0D, 0x1B, 0x2A)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def set_slide_bg(slide, color=BG_DARK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    font_size=18,
    color=LIGHT_GRAY,
    bold=False,
    alignment=PP_ALIGN.LEFT,
):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Calibri"
    p.alignment = alignment
    return txBox


def add_rich_textbox(slide, left, top, width, height, runs):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    for text, size, color, bold in runs:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = "Calibri"
    return txBox


def add_bullet_list(
    slide, left, top, width, items, font_size=20, bullet_color=ACCENT_RED
):
    txBox = slide.shapes.add_textbox(left, top, width, Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(10)
        p.space_after = Pt(6)

        if " — " in item:
            bold_part, rest = item.split(" — ", 1)
            run_b = p.add_run()
            run_b.text = bold_part + " — "
            run_b.font.size = Pt(font_size)
            run_b.font.color.rgb = WHITE
            run_b.font.bold = True
            run_b.font.name = "Calibri"
            run_r = p.add_run()
            run_r.text = rest
            run_r.font.size = Pt(font_size)
            run_r.font.color.rgb = LIGHT_GRAY
            run_r.font.name = "Calibri"
        else:
            run = p.add_run()
            run.text = item
            run.font.size = Pt(font_size)
            run.font.color.rgb = LIGHT_GRAY
            run.font.name = "Calibri"

        pPr = p._p.get_or_add_pPr()
        for old in pPr.findall(qn("a:buNone")):
            pPr.remove(old)
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        buClr = pPr.makeelement(qn("a:buClr"), {})
        buClr.append(buClr.makeelement(qn("a:srgbClr"), {"val": str(bullet_color)}))
        pPr.append(buClr)
        pPr.append(pPr.makeelement(qn("a:buSzPct"), {"val": "120000"}))

    return txBox


def add_divider(slide, top, color=ACCENT_RED):
    shape = slide.shapes.add_shape(1, Inches(0.8), top, Inches(2), Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_card(
    slide, left, top, width, height, title, body_items, accent_color=ACCENT_RED
):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG_MID
    shape.line.color.rgb = accent_color
    shape.line.width = Pt(1.5)

    add_textbox(
        slide,
        left + Inches(0.25),
        top + Inches(0.15),
        width - Inches(0.5),
        Inches(0.4),
        title,
        font_size=18,
        color=accent_color,
        bold=True,
    )

    y = top + Inches(0.55)
    for item in body_items:
        add_textbox(
            slide,
            left + Inches(0.25),
            y,
            width - Inches(0.5),
            Inches(0.35),
            item,
            font_size=14,
            color=LIGHT_GRAY,
        )
        y += Inches(0.32)


def add_table(slide, left, top, width, rows_data, col_widths_pct):
    n_rows = len(rows_data)
    n_cols = len(rows_data[0])
    table_h = Inches(0.45 * n_rows)
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, table_h)
    table = shape.table

    for i, pct in enumerate(col_widths_pct):
        table.columns[i].width = int(width * pct)

    for r_idx, row_data in enumerate(rows_data):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = cell_text
            run.font.name = "Calibri"

            if r_idx == 0:
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = ACCENT_RED
            else:
                run.font.size = Pt(13)
                run.font.color.rgb = LIGHT_GRAY
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_BG

    return shape


# ============================================================
# SLIDE 1 — Title
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(2.0),
    Inches(11),
    Inches(1),
    "Z-Stream Test Generator",
    font_size=44,
    color=WHITE,
    bold=True,
    alignment=PP_ALIGN.LEFT,
)

add_divider(slide, Inches(3.1))

add_textbox(
    slide,
    Inches(0.8),
    Inches(3.4),
    Inches(10),
    Inches(0.6),
    "AI-Driven Verification Test Generation for ODF Bug Fixes",
    font_size=24,
    color=ACCENT_BLUE,
    alignment=PP_ALIGN.LEFT,
)

add_textbox(
    slide,
    Inches(0.8),
    Inches(4.3),
    Inches(10),
    Inches(0.5),
    "ODF QE  |  September 2026",
    font_size=18,
    color=DIM_GRAY,
    alignment=PP_ALIGN.LEFT,
)


# ============================================================
# SLIDE 2 — The Problem
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "The Problem",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

add_bullet_list(
    slide,
    Inches(0.8),
    Inches(1.6),
    Inches(11.5),
    [
        "Repetitive manual work — Every z-stream fix requires a hand-written "
        "verification test following strict ocs-ci conventions",
        "Backport multiplication — The same fix is backported across 3-5 ODF "
        "versions, each needing separate verification tracking",
        "One-time verification gap — Manually verified fixes produce no lasting "
        "automation artifact; regressions can silently reappear",
        "Scaling bottleneck — A typical z-stream release has 15-25 bugs; "
        "writing tests for all of them competes with higher-value engineering work",
        "Convention drift — Manual test writing leads to inconsistent patterns "
        "across the test suite (logging, cleanup, markers, base classes)",
    ],
    font_size=20,
)


# ============================================================
# SLIDE 3 — The Solution (Overview)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "The Solution",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

add_textbox(
    slide,
    Inches(0.8),
    Inches(1.5),
    Inches(11),
    Inches(1.0),
    "An end-to-end pipeline that reads Jira bugs, understands the upstream fix, "
    "and generates production-ready ocs-ci pytest tests — complete with helper "
    "functions, draft PRs, and backport branches.",
    font_size=20,
    color=LIGHT_GRAY,
)

add_textbox(
    slide,
    Inches(0.8),
    Inches(2.5),
    Inches(11),
    Inches(0.5),
    "One command processes an entire z-stream release:",
    font_size=18,
    color=DIM_GRAY,
)

add_textbox(
    slide,
    Inches(1.2),
    Inches(3.0),
    Inches(10),
    Inches(0.5),
    "python -m ocs_ci.utility.zstream_test_gen --fix-version odf-4.22.5 --confidence high",
    font_size=18,
    color=ACCENT_GREEN,
    bold=True,
)

# Key numbers
cards = [
    (Inches(0.8), "~15 min", "End-to-end for a full\nz-stream release"),
    (Inches(4.0), "15-25", "Bugs processed\nper release"),
    (Inches(7.2), "1 test", "Per unique fix across\nall backported versions"),
    (Inches(10.4), "Draft PRs", "Auto-created with\nbackports to each branch"),
]
for left, number, desc in cards:
    add_card(
        slide,
        left,
        Inches(4.0),
        Inches(2.7),
        Inches(2.2),
        number,
        desc.split("\n"),
        accent_color=ACCENT_BLUE,
    )


# ============================================================
# SLIDE 4 — Pipeline Stages
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "Pipeline Stages",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

stages = [
    ("0a", "Fix Reviews", "Auto-fix CodeRabbit comments on open PRs", DIM_GRAY),
    ("0b", "Learn", "Collect feedback from merged/closed PRs", DIM_GRAY),
    ("1", "Collect", "Query Jira for all bugs in the fix version", ACCENT_BLUE),
    ("2", "Enrich", "Fetch upstream PR diffs + AI root cause analysis", ACCENT_BLUE),
    ("3", "Classify", "Automatable / DR / Manual-only / Already covered", ACCENT_BLUE),
    ("4", "Deduplicate", "Group backported clones, keep richest context", ACCENT_RED),
    (
        "5",
        "Skip Open PRs",
        "Skip bugs that already have PRs from prior runs",
        ACCENT_RED,
    ),
    ("6", "Score", "Confidence scoring (7 signals, 0-100%)", ACCENT_RED),
    ("7", "Generate", "AI test generation with similar-test examples", ACCENT_GREEN),
    ("8", "Extract", "Second AI pass: refactor into helper functions", ACCENT_GREEN),
    ("9", "Validate", "Syntax, imports, patterns, flake8, secrets scan", ACCENT_YELLOW),
    (
        "10",
        "Publish",
        "Draft PRs + backports + Jira labels + secrets gate",
        ACCENT_ORANGE,
    ),
]

col_x = [Inches(0.5), Inches(3.7), Inches(6.9), Inches(10.1)]
for i, (num, name, desc, color) in enumerate(stages):
    col = i % 4
    row = i // 4
    x = col_x[col]
    y = Inches(1.6) + Inches(row * 1.9)

    shape = slide.shapes.add_shape(1, x, y, Inches(2.9), Inches(1.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG_MID
    shape.line.color.rgb = color
    shape.line.width = Pt(2)

    add_rich_textbox(
        slide,
        x + Inches(0.15),
        y + Inches(0.12),
        Inches(2.6),
        Inches(0.4),
        [(num + ". ", 14, color, True), (name, 15, WHITE, True)],
    )

    add_textbox(
        slide,
        x + Inches(0.15),
        y + Inches(0.6),
        Inches(2.6),
        Inches(0.8),
        desc,
        font_size=12,
        color=LIGHT_GRAY,
    )


# ============================================================
# SLIDE 5 — Confidence Scoring
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "Confidence Scoring",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

add_textbox(
    slide,
    Inches(0.8),
    Inches(1.4),
    Inches(11),
    Inches(0.6),
    "Each bug is scored on 7 signals to determine how likely the generated "
    "test is to be correct and useful.",
    font_size=18,
    color=LIGHT_GRAY,
)

add_table(
    slide,
    Inches(0.8),
    Inches(2.2),
    Inches(7),
    [
        ["Signal", "Weight", "Why It Matters"],
        ["Has upstream PR URL", "+2", "Links test to the actual fix"],
        ["Has upstream PR diff", "+2", "AI can read the code changes"],
        ["Substantive description", "+1", "Bug context for test logic"],
        ["Reproduction steps", "+1", "Concrete verification path"],
        ["AI verification steps", "+2", "Quality of generated spec"],
        ["Similar existing tests", "+1", "Patterns to follow"],
        ["Known component", "+1", "Correct directory + base class"],
    ],
    [0.35, 0.15, 0.50],
)

# Level mapping on the right
add_card(
    slide,
    Inches(8.5),
    Inches(2.2),
    Inches(4),
    Inches(1.8),
    "Confidence Levels",
    [
        "High: >= 70% (7+/10)",
        "Medium: >= 40% (4+/10)",
        "Low: < 40% (0-3/10)",
    ],
    accent_color=ACCENT_GREEN,
)

add_card(
    slide,
    Inches(8.5),
    Inches(4.4),
    Inches(4),
    Inches(1.8),
    "CLI Filter",
    [
        "--confidence high",
        "Only publish PRs for tests at",
        "or above this threshold",
    ],
    accent_color=ACCENT_YELLOW,
)


# ============================================================
# SLIDE 6 — What Gets Generated
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "What Gets Generated",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

# Left column — test file
add_textbox(
    slide,
    Inches(0.8),
    Inches(1.5),
    Inches(5),
    Inches(0.4),
    "Test File",
    font_size=22,
    color=ACCENT_BLUE,
    bold=True,
)
add_bullet_list(
    slide,
    Inches(0.8),
    Inches(2.0),
    Inches(5.5),
    [
        "Correct base class — ManageTest, E2ETest, etc.",
        "Squad + tier markers — green_squad, tier1, etc.",
        "Jira traceability — @jira_ticket('DFBUGS-XXXX')",
        "Fixture-based cleanup — pvc_factory, pod_factory",
        "Custom logging — test_step, assertion log levels",
        "Version gating — skipif_ocs_version decorators",
        "Google-style docstrings — with bug context",
    ],
    font_size=16,
)

# Right column — helpers + PR
add_textbox(
    slide,
    Inches(7),
    Inches(1.5),
    Inches(5),
    Inches(0.4),
    "Helper Functions",
    font_size=22,
    color=ACCENT_GREEN,
    bold=True,
)
add_bullet_list(
    slide,
    Inches(7),
    Inches(2.0),
    Inches(5.5),
    [
        "Second AI pass — identifies inline patterns to extract",
        "Targets correct module — e.g., ocs_ci/helpers/helpers.py",
        "Staged for review — placed under zstream_helpers/",
        "Non-blocking — falls back to original if extraction fails",
    ],
    font_size=16,
    bullet_color=ACCENT_GREEN,
)

add_textbox(
    slide,
    Inches(7),
    Inches(3.9),
    Inches(5),
    Inches(0.4),
    "Draft PR",
    font_size=22,
    color=ACCENT_ORANGE,
    bold=True,
)
add_bullet_list(
    slide,
    Inches(7),
    Inches(4.4),
    Inches(5.5),
    [
        "Test + helpers — committed together in one PR",
        "Structured body — summary, root cause, execution context",
        "Review checklist — built into the PR description",
        "Backport PRs — auto-created for each release branch",
        "Jira updated — comment with PR links + tracking label",
    ],
    font_size=16,
    bullet_color=ACCENT_ORANGE,
)


# ============================================================
# SLIDE 7 — Benefits
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "Benefits",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

benefits = [
    (
        "Reduce QE overhead",
        "Engineers review and refine instead of writing boilerplate "
        "verification tests from scratch for every z-stream fix",
        ACCENT_RED,
    ),
    (
        "Increase automation coverage",
        "Deduplication across backports + automatic release branch PRs "
        "ensure every version gets coverage without extra effort",
        ACCENT_BLUE,
    ),
    (
        "Permanent regression tests",
        "Every verified fix becomes part of the regression suite, "
        "catching silent regressions in future releases automatically",
        ACCENT_GREEN,
    ),
    (
        "Self-improving",
        "Learns from CodeRabbit reviews and human reviewer corrections "
        "— the same mistakes are never repeated in future generations",
        ACCENT_YELLOW,
    ),
    (
        "Consistent conventions",
        "Generated tests follow ocs-ci patterns by design: base classes, "
        "markers, fixtures, logging, cleanup — no review cycles for style",
        ACCENT_ORANGE,
    ),
    (
        "Safe by default",
        "Two-layer secrets scanning (validation + publish gate) prevents "
        "credentials, keys, and tokens from reaching GitHub",
        ACCENT_BLUE,
    ),
    (
        "Retroactive coverage",
        "Process past fix versions to generate tests for bugs that were "
        "only verified manually — systematically close automation gaps",
        ACCENT_RED,
    ),
    (
        "Faster qualification",
        "Front-loads test generation so verification tests are ready for "
        "review as soon as fixes are merged, compressing release timelines",
        ACCENT_GREEN,
    ),
]

for i, (title, desc, color) in enumerate(benefits):
    col = i % 4
    row = i // 4
    x = Inches(0.3) + Inches(col * 3.25)
    y = Inches(1.5) + Inches(row * 2.8)
    add_card(slide, x, y, Inches(3.0), Inches(2.4), title, [desc], accent_color=color)


# ============================================================
# SLIDE 8 — Architecture
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "Architecture",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

components = [
    (
        "Jira Cloud",
        "redhat.atlassian.net\nDFBUGS project\nBug data + comments + clone chains",
        ACCENT_RED,
        Inches(0.5),
    ),
    (
        "GitHub API",
        "Upstream PR diffs\nFork-based PR creation\nGit Data API (serverless)",
        ACCENT_BLUE,
        Inches(3.5),
    ),
    (
        "Claude AI",
        "Vertex AI (Claude Opus)\nBug analysis + test generation\nHelper extraction",
        ACCENT_GREEN,
        Inches(6.5),
    ),
    (
        "ocs-ci",
        "Codebase conventions\nSimilar test discovery\nValidation (ast, flake8, pytest)",
        ACCENT_YELLOW,
        Inches(9.5),
    ),
]

for title, desc, color, x in components:
    shape = slide.shapes.add_shape(1, x, Inches(1.6), Inches(3), Inches(2.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG_MID
    shape.line.color.rgb = color
    shape.line.width = Pt(2)

    add_textbox(
        slide,
        x + Inches(0.2),
        Inches(1.7),
        Inches(2.6),
        Inches(0.35),
        title,
        font_size=18,
        color=color,
        bold=True,
        alignment=PP_ALIGN.CENTER,
    )
    for j, line in enumerate(desc.split("\n")):
        add_textbox(
            slide,
            x + Inches(0.2),
            Inches(2.2) + Inches(j * 0.35),
            Inches(2.6),
            Inches(0.3),
            line,
            font_size=13,
            color=LIGHT_GRAY,
            alignment=PP_ALIGN.CENTER,
        )

# Flow arrows (text-based)
for x in [Inches(3.2), Inches(6.2), Inches(9.2)]:
    add_textbox(
        slide,
        x,
        Inches(2.3),
        Inches(0.5),
        Inches(0.5),
        "→",
        font_size=28,
        color=DIM_GRAY,
        alignment=PP_ALIGN.CENTER,
    )

# Module listing
add_textbox(
    slide,
    Inches(0.8),
    Inches(4.2),
    Inches(11),
    Inches(0.4),
    "Module Structure",
    font_size=20,
    color=WHITE,
    bold=True,
)

modules = [
    ["Module", "Responsibility"],
    [
        "cli.py",
        "CLI entry point, argument parsing (--maintain, --fix-reviews, --learn-only)",
    ],
    ["config.py", "YAML + env var configuration loading"],
    ["pipeline.py", "Orchestrates all stages end-to-end incl. feedback loop"],
    ["jira_client.py", "Jira queries, parsing, clone chain walking"],
    [
        "github_client.py",
        "PR diffs, fork-based PR creation, CodeRabbit comments, Git Data API",
    ],
    ["generator.py", "Claude AI calls with auto-injected correction rules"],
    ["classifier.py", "Bug classification + confidence scoring"],
    [
        "validator.py",
        "Syntax, import, pattern, flake8, secrets scanning, pytest validation",
    ],
    ["publisher.py", "PR publishing, backports, Jira labeling, final secrets gate"],
    ["feedback.py", "Automated feedback loop: ReviewFixer + FeedbackCollector"],
    ["prompts.py", "All AI prompt templates"],
    ["models.py", "Data models (BugInfo, GeneratedTest, HelperSpec, etc.)"],
]

add_table(slide, Inches(0.8), Inches(4.7), Inches(11.5), modules, [0.22, 0.78])


# ============================================================
# SLIDE 9 — Results (odf-4.22.5)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "Validation Run: odf-4.22.5",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

# Stats cards
stats = [
    ("23", "Bugs\ncollected", ACCENT_BLUE),
    ("15", "Tests\ngenerated", ACCENT_GREEN),
    ("8", "Manual-only\n(UI/perf/CVE)", ACCENT_YELLOW),
    ("100%", "Validation\npass rate", ACCENT_GREEN),
    ("4", "High-confidence\nPRs ready", ACCENT_RED),
]
for i, (num, label, color) in enumerate(stats):
    x = Inches(0.5) + Inches(i * 2.6)
    shape = slide.shapes.add_shape(1, x, Inches(1.6), Inches(2.2), Inches(2.0))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG_MID
    shape.line.color.rgb = color
    shape.line.width = Pt(2)

    add_textbox(
        slide,
        x,
        Inches(1.7),
        Inches(2.2),
        Inches(0.8),
        num,
        font_size=40,
        color=color,
        bold=True,
        alignment=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        x,
        Inches(2.6),
        Inches(2.2),
        Inches(0.8),
        label,
        font_size=16,
        color=LIGHT_GRAY,
        alignment=PP_ALIGN.CENTER,
    )

add_textbox(
    slide,
    Inches(0.8),
    Inches(4.0),
    Inches(11),
    Inches(0.4),
    "What the pipeline produced:",
    font_size=20,
    color=WHITE,
    bold=True,
)

add_bullet_list(
    slide,
    Inches(0.8),
    Inches(4.5),
    Inches(11),
    [
        "15 complete pytest files — following ocs-ci conventions with proper markers, "
        "fixtures, and cleanup",
        "2-6 helper functions per test — extracted and staged for target module integration",
        "Automatic deduplication — backported clones grouped, one test per unique fix",
        "Confidence distribution — 4 high, 7 medium, 4 low; CLI filter controls which get PRs",
    ],
    font_size=18,
)


# ============================================================
# SLIDE 10 — Automated Feedback Loop
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "Automated Feedback Loop",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

add_textbox(
    slide,
    Inches(0.8),
    Inches(1.4),
    Inches(11),
    Inches(0.6),
    "The tool learns from every PR interaction — CodeRabbit reviews, "
    "human corrections, and merge outcomes — and applies those lessons "
    "to all future generations automatically.",
    font_size=18,
    color=LIGHT_GRAY,
)

# Left card — CodeRabbit auto-fix
add_card(
    slide,
    Inches(0.5),
    Inches(2.3),
    Inches(3.8),
    Inches(3.5),
    "CodeRabbit Review Fixing",
    [
        "Detects new CodeRabbit comments on open PRs",
        "Reads comments + current file contents",
        "Generates fixes via Claude with ocs-ci conventions",
        "Pushes fix commit directly to the PR branch",
        "Tracks comment count to detect new reviews",
        "CLI: --fix-reviews or --maintain",
    ],
    accent_color=ACCENT_BLUE,
)

# Center card — merged PR learning
add_card(
    slide,
    Inches(4.8),
    Inches(2.3),
    Inches(3.8),
    Inches(3.5),
    "Merged PR Learning",
    [
        "Finds merged PRs with ai-generated label",
        "Compares original vs final merged code",
        "Extracts reusable correction rules via Claude",
        "Rules stored in ~/.zstream_test_gen_corrections.yaml",
        "Auto-injected into system prompt for future runs",
        "CLI: --learn-only or --maintain",
    ],
    accent_color=ACCENT_GREEN,
)

# Right card — secrets scanning
add_card(
    slide,
    Inches(9.1),
    Inches(2.3),
    Inches(3.8),
    Inches(3.5),
    "Secrets Protection",
    [
        "10 regex patterns for credentials, keys, tokens",
        "Scans during validation (blocking check)",
        "Final gate at publish: all files + PR body",
        "Catches private keys, pull secrets, AWS keys",
        "JWTs, Bearer tokens, registry credentials",
        "PR is blocked if any match is found",
    ],
    accent_color=ACCENT_RED,
)

# Bottom — maintenance mode
add_textbox(
    slide,
    Inches(0.8),
    Inches(6.2),
    Inches(11),
    Inches(0.4),
    "Maintenance mode (--maintain) runs review fixing + feedback "
    "collection without generating new tests — ideal for periodic cron jobs.",
    font_size=16,
    color=DIM_GRAY,
)


# ============================================================
# SLIDE 11 — Future & How to Use
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_textbox(
    slide,
    Inches(0.8),
    Inches(0.5),
    Inches(10),
    Inches(0.6),
    "Getting Started & What's Next",
    font_size=36,
    color=WHITE,
    bold=True,
)
add_divider(slide, Inches(1.15))

# Left — how to use
add_textbox(
    slide,
    Inches(0.8),
    Inches(1.5),
    Inches(5),
    Inches(0.4),
    "How to Use",
    font_size=24,
    color=ACCENT_BLUE,
    bold=True,
)

add_bullet_list(
    slide,
    Inches(0.8),
    Inches(2.1),
    Inches(5.5),
    [
        "Configure — Run --init-config, add Jira/GitHub/Claude credentials",
        "Fork — Fork red-hat-storage/ocs-ci, set fork_repo in config",
        "Generate — Run with --fix-version and --confidence level",
        "Maintain — Run --maintain to fix reviews + learn from merged PRs",
        "Fix reviews — Run --fix-reviews to address CodeRabbit comments",
        "Learn only — Run --learn-only to collect feedback from closed PRs",
        "Single bug — Use --bug DFBUGS-XXXX for individual bugs",
    ],
    font_size=17,
)

# Right — what's next
add_textbox(
    slide,
    Inches(7),
    Inches(1.5),
    Inches(5),
    Inches(0.4),
    "What's Next",
    font_size=24,
    color=ACCENT_GREEN,
    bold=True,
)

add_bullet_list(
    slide,
    Inches(7),
    Inches(2.1),
    Inches(5.5),
    [
        'Improve component detection — reduce "unknown" classifications',
        "CI integration — trigger automatically on new fix versions",
        "Historical backfill — process past releases for coverage gaps",
        "Team adoption — shared credentials + CI-based workflow",
        "Scheduled maintenance — cron-based --maintain runs",
    ],
    font_size=17,
    bullet_color=ACCENT_GREEN,
)

# Bottom — key links
add_textbox(
    slide,
    Inches(0.8),
    Inches(5.5),
    Inches(11),
    Inches(0.4),
    "Location: ocs_ci/utility/zstream_test_gen/",
    font_size=18,
    color=DIM_GRAY,
)
add_textbox(
    slide,
    Inches(0.8),
    Inches(5.9),
    Inches(11),
    Inches(0.4),
    "Config: ~/.zstream_test_gen.yaml   |   Output: .zstream_gen_output/",
    font_size=18,
    color=DIM_GRAY,
)


# ============================================================
# Save
# ============================================================
out_path = "zstream_test_gen_overview.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
